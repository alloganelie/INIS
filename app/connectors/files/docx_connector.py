"""DOCX connector for INIS per §9.1."""

from pathlib import Path

from app.connectors.base import (
    ConnectorMetadata,
    HealthStatus,
    Query,
    RawSource,
    SourceCandidate,
    SourceConnector,
    SourceMetadata,
)

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _read_docx(data: bytes) -> tuple[list[str], int, list[str]]:
    """Extract paragraphs, table count and style names via python-docx.

    Args:
        data: Raw .docx bytes.

    Returns:
        Tuple of (non-empty paragraph texts, table count, sorted style names).
    """
    import io

    import docx

    document = docx.Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    styles = sorted({p.style.name for p in document.paragraphs if p.style is not None})
    return paragraphs, len(document.tables), styles


class DOCXConnector:
    """DOCX file connector (python-docx for text + structure)."""

    connector_id: str = "docx-connector"
    supported_source_types: list[str] = ["docx"]

    def __init__(self, base_path: str = ".") -> None:
        """Initialize the DOCX connector.

        Args:
            base_path: Base directory for DOCX file discovery.
        """
        self._base_path = Path(base_path)

    async def discover(self, query: Query) -> list[SourceCandidate]:
        """Discover DOCX files matching the query.

        Args:
            query: Query for source discovery.

        Returns:
            List of source candidates.
        """
        candidates: list[SourceCandidate] = []
        for docx_file in self._base_path.glob("*.docx"):
            if query.query_string.lower() in docx_file.name.lower():
                candidates.append(
                    SourceCandidate(
                        source_id=f"docx-{docx_file.stem}",
                        location=str(docx_file),
                        metadata={"filename": docx_file.name},
                    )
                )
        return candidates

    async def retrieve(self, candidate: SourceCandidate) -> RawSource:
        """Retrieve raw DOCX bytes from a candidate.

        Args:
            candidate: Source candidate to retrieve.

        Returns:
            Raw source data.
        """
        file_path = Path(candidate.location)
        content = file_path.read_bytes()
        return RawSource(
            source_id=candidate.source_id,
            data=content,
            content_type=CONTENT_TYPE,
            metadata=candidate.metadata,
        )

    async def inspect(self, raw: RawSource) -> SourceMetadata:
        """Inspect raw DOCX: size, paragraph/table counts and styles.

        Args:
            raw: Raw source data.

        Returns:
            Source metadata (size-only when python-docx is unavailable).
        """
        payload = raw.data if isinstance(raw.data, bytes) else raw.data.encode("utf-8")
        try:
            paragraphs, table_count, styles = _read_docx(bytes(payload))
        except ImportError:
            return SourceMetadata(
                source_id=raw.source_id,
                size_bytes=len(payload),
                record_count=0,
            )
        except Exception:
            return SourceMetadata(
                source_id=raw.source_id,
                size_bytes=len(payload),
                record_count=0,
            )
        schema = {
            "paragraphs": str(len(paragraphs)),
            "tables": str(table_count),
        }
        if styles:
            schema["styles"] = ",".join(styles)
        return SourceMetadata(
            source_id=raw.source_id,
            size_bytes=len(payload),
            record_count=len(paragraphs),
            schema=schema,
        )

    async def health_check(self) -> HealthStatus:
        """Check if the connector is healthy (python-docx importable).

        Returns:
            Health status.
        """
        try:
            import docx  # noqa: F401
        except ImportError:
            return HealthStatus(healthy=False, message="DOCX connector missing python-docx")
        return HealthStatus(healthy=True, message="DOCX connector healthy")

    async def metadata(self) -> ConnectorMetadata:
        """Get connector metadata.

        Returns:
            Connector metadata.
        """
        return ConnectorMetadata(
            connector_id=self.connector_id,
            name="DOCX Connector",
            version="1.0.0",
            supported_source_types=self.supported_source_types,
        )
