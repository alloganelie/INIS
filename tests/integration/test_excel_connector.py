"""§9.1 — the Excel connector on real ``.xlsx`` workbooks.

Excel is the format of the datasets INIS imports, and it was implemented without
a test (`docs/SPEC_COVERAGE.md` §9). These tests build real workbooks with
``openpyxl`` in a temporary directory and walk the connector contract: discovery,
retrieval, inspection, health and identity.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.connectors.base import Query, RawSource, SourceCandidate
from app.connectors.files.excel_connector import ExcelConnector

pytestmark = pytest.mark.asyncio

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _write_xlsx(path: Path, rows: list[tuple[str, int]]) -> Path:
    """Write a real workbook containing *rows*."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(("city", "population"))
    for city, population in rows:
        sheet.append((city, population))
    workbook.save(path)
    return path


@pytest.fixture
def excel_dir(tmp_path: Path) -> Path:
    """Return a directory holding two real workbooks."""
    _write_xlsx(tmp_path / "villes_population.xlsx", [("Paris", 2145906), ("Berlin", 3645000)])
    _write_xlsx(tmp_path / "regions_population.xlsx", [("IDF", 12262702)])
    return tmp_path


class TestDiscovery:
    """§9.1 — the connector finds the workbooks the query names."""

    async def test_only_matching_files_are_discovered(self, excel_dir: Path) -> None:
        """Discovery filters on the file name and returns candidates."""
        connector = ExcelConnector(str(excel_dir))

        candidates = await connector.discover(Query(query_string="villes"))

        assert [candidate.source_id for candidate in candidates] == ["excel-villes_population"]

    async def test_no_match_yields_no_candidate(self, excel_dir: Path) -> None:
        """An unmatched query returns nothing rather than every workbook."""
        connector = ExcelConnector(str(excel_dir))

        assert await connector.discover(Query(query_string="sans-rapport")) == []


class TestRetrieval:
    """§9.1 — the workbook is retrieved byte-identical."""

    async def test_retrieve_returns_the_workbook_bytes(self, excel_dir: Path) -> None:
        """The retrieved payload is the file on disk, with the OOXML MIME type."""
        connector = ExcelConnector(str(excel_dir))
        candidate = (await connector.discover(Query(query_string="regions")))[0]

        raw = await connector.retrieve(candidate)

        assert isinstance(raw, RawSource)
        assert raw.content_type == CONTENT_TYPE
        assert raw.data == (excel_dir / "regions_population.xlsx").read_bytes()

    async def test_retrieved_bytes_are_a_readable_workbook(self, excel_dir: Path) -> None:
        """What was transported is still a valid workbook (integrity, §18.1)."""
        import io

        connector = ExcelConnector(str(excel_dir))
        candidate = (await connector.discover(Query(query_string="villes")))[0]

        raw = await connector.retrieve(candidate)

        workbook = openpyxl.load_workbook(io.BytesIO(raw.data))
        sheet = workbook.active
        assert sheet.max_row == 3
        assert sheet["A2"].value == "Paris"


class TestInspection:
    """§9.1 — inspection reports what it measured and nothing more."""

    async def test_inspection_reports_the_size_only(self, excel_dir: Path) -> None:
        """``record_count`` stays ``None``: the stub does not count rows."""
        connector = ExcelConnector(str(excel_dir))
        candidate = (await connector.discover(Query(query_string="villes")))[0]

        metadata = await connector.inspect(await connector.retrieve(candidate))

        assert metadata.size_bytes == (excel_dir / "villes_population.xlsx").stat().st_size
        assert metadata.record_count is None

    async def test_inspection_accepts_a_text_payload(self) -> None:
        """A ``str`` payload is measured by its UTF-8 length, not crashed on."""
        connector = ExcelConnector(".")

        metadata = await connector.inspect(
            RawSource(
                source_id="excel-text",
                data="a,b\n1,2",
                content_type=CONTENT_TYPE,
                metadata={},
            )
        )

        assert metadata.size_bytes == len("a,b\n1,2".encode("utf-8"))

    async def test_missing_file_surfaces_the_underlying_error(self, tmp_path: Path) -> None:
        """A candidate pointing at a missing path fails loudly."""
        connector = ExcelConnector(str(tmp_path))
        candidate = SourceCandidate(
            source_id="excel-ghost", location=str(tmp_path / "ghost.xlsx"), metadata={}
        )

        with pytest.raises(FileNotFoundError):
            await connector.retrieve(candidate)


class TestConnectorContract:
    """§9 — the connector exposes the same surface as every other one."""

    async def test_health_is_reported(self) -> None:
        """A stub connector is still healthy or explicitly not."""
        health = await ExcelConnector().health_check()

        assert health.healthy is True

    async def test_metadata_declares_both_excel_types(self) -> None:
        """``xlsx`` and the legacy ``xls`` are both advertised."""
        metadata = await ExcelConnector().metadata()

        assert metadata.connector_id == "excel-connector"
        assert metadata.supported_source_types == ["xlsx", "xls"]

