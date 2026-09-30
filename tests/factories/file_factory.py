"""Deterministic *file* builders for the §9.1 ingestion tests (§33.2).

The ingestion is exercised on real bytes, never on a mock reader: a PDF whose
pages carry distinct sentences, a DOCX with paragraphs and a table, a PNG with an
embedded ``ImageDescription``, a workbook with several sheets. Each builder is
pure and writes nothing: it returns the bytes an upload would carry, so a test
can prove what INIS extracts from *that* document.

A minimal PDF is built by hand because no writer engine is installed (§4.1, C22):
the objects, the cross-reference table and the offsets are computed, which is
exactly what makes the fixture a valid document rather than a plausible prefix.
"""

from __future__ import annotations

from collections.abc import Sequence
from io import BytesIO

__all__ = [
    "docx_bytes",
    "pdf_bytes",
    "png_bytes",
    "workbook_bytes",
]


def pdf_bytes(pages: Sequence[str]) -> bytes:
    """Return a valid PDF holding one text line per entry of *pages*.

    Args:
        pages: The text of each page, in page order. An empty string yields a
            page with no text at all — the scanned-page case (§9.2).

    Returns:
        The PDF bytes, with a real cross-reference table.
    """
    kids = b" ".join(b"%d 0 R" % (4 + index * 2) for index in range(len(pages)))
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(pages)),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    for index, text in enumerate(pages):
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 24 Tf 100 700 Td ({safe}) Tj ET".encode("latin-1")
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 3 0 R >> >> >>" % (5 + index * 2)
        )
        objects.append(
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += (
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, xref_at)
    )
    return bytes(out)


def docx_bytes(
    paragraphs: Sequence[str],
    *,
    tables: Sequence[Sequence[Sequence[str]]] = (),
) -> bytes:
    """Return a ``.docx`` holding *paragraphs* (empty ones included) and *tables*.

    Args:
        paragraphs: Body paragraphs, in order; an empty string is a blank
            paragraph, which keeps the paragraph numbering of the document.
        tables: One entry per table, each a sequence of rows of cell texts.
    """
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for rows in tables:
        table = document.add_table(rows=len(rows), cols=len(rows[0]) if rows else 1)
        for row_index, row in enumerate(rows):
            for cell_index, cell in enumerate(row):
                table.cell(row_index, cell_index).text = cell
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def png_bytes(text: str | None = None, *, size: tuple[int, int] = (8, 6)) -> bytes:
    """Return a PNG, optionally carrying *text* as its ``ImageDescription`` tag.

    The tag is the only textual content INIS may quote from an image: there is no
    OCR and no captioning in V1 (§9.2).
    """
    from PIL import Image
    from PIL.PngImagePlugin import PngInfo

    image = Image.new("RGB", size, color=(10, 20, 30))
    buffer = BytesIO()
    if text is None:
        image.save(buffer, format="PNG")
    else:
        info = PngInfo()
        info.add_text("ImageDescription", text)
        image.save(buffer, format="PNG", pnginfo=info)
    return buffer.getvalue()


def workbook_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    """Return an ``.xlsx`` workbook with one worksheet per ``{name: rows}`` entry.

    Args:
        sheets: Worksheet name → rows (the first row is the header, as read by
            :func:`app.tools.files.excel_reader.load_excel`).
    """
    from openpyxl import Workbook

    workbook = Workbook()
    default = workbook.active
    first = True
    for name, rows in sheets.items():
        worksheet = default if first else workbook.create_sheet()
        worksheet.title = name
        for row in rows:
            worksheet.append(row)
        first = False
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
