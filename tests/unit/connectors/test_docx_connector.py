"""Tests for DOCXConnector per §9.1 (tmp_path, python-docx fixture)."""

from pathlib import Path

import pytest

from app.connectors.base import Query
from app.connectors.files.docx_connector import DOCXConnector


def _make_docx(path: Path) -> None:
    """Build a minimal .docx: 2 paragraphs + 1 table."""
    import docx

    document = docx.Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "cell-a"
    document.save(str(path))


@pytest.mark.asyncio
async def test_docx_discover(tmp_path: Path) -> None:
    """DOCX files matching the query are discovered."""
    _make_docx(tmp_path / "report_test.docx")

    candidates = await DOCXConnector(base_path=str(tmp_path)).discover(
        Query(query_string="report")
    )

    assert len(candidates) == 1
    assert candidates[0].source_id == "docx-report_test"


@pytest.mark.asyncio
async def test_docx_retrieve(tmp_path: Path) -> None:
    """Raw DOCX bytes are retrieved with the OOXML content type."""
    _make_docx(tmp_path / "report_test.docx")
    connector = DOCXConnector(base_path=str(tmp_path))

    candidates = await connector.discover(Query(query_string="report"))
    raw = await connector.retrieve(candidates[0])

    assert raw.content_type.startswith("application/vnd.openxmlformats")
    assert isinstance(raw.data, bytes)
    assert raw.data[:2] == b"PK"


@pytest.mark.asyncio
async def test_docx_inspect(tmp_path: Path) -> None:
    """Inspect reports paragraph/table counts and styles."""
    _make_docx(tmp_path / "report_test.docx")
    connector = DOCXConnector(base_path=str(tmp_path))

    candidates = await connector.discover(Query(query_string="report"))
    raw = await connector.retrieve(candidates[0])
    metadata = await connector.inspect(raw)

    assert metadata.record_count == 2
    assert metadata.schema is not None
    assert metadata.schema["tables"] == "1"
    assert "Normal" in metadata.schema["styles"]
