"""§24.2/§24.3 — the delivered files are generated, deterministic and honest.

The point of these tests is not that "a file is produced": it is that the bytes
are reproducible (so the ``sha256`` of the §24.2 record can be recomputed by the
client that received the file), that nothing is invented (§37) and that a format
the V1 stack cannot produce is *refused* rather than faked (§4.1).
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from lxml import etree
from openpyxl import load_workbook

from app.artifacts.generators import (
    AVAILABLE_FORMATS,
    FILE_FORMATS,
    generate_artifact,
    generate_csv,
    generate_json,
    generate_xlsx,
    generate_xml,
    pdf_supported,
)
from app.artifacts.generators.csv_generator import ENCODING
from app.artifacts.generators.xlsx_generator import SHEET_ORDER, SHEET_TITLES
from app.core.errors import ValidationError

ROWS = [
    {"information_id": "INF_1", "text": "Paris", "confidence_score": 0.9},
    {
        "information_id": "INF_2",
        "text": "Lyon",
        "confidence_score": None,
        "provenance": {"url": "https://exemple.fr"},
    },
]


class TestCsvExport:
    """§24.3 « données demandées » — the flat export keeps its promises."""

    def test_header_is_the_caller_order(self) -> None:
        """A requested column order is honoured, not re-sorted."""
        document = generate_csv(ROWS, columns=["text", "information_id"]).decode(ENCODING)

        assert document.splitlines()[0] == "text;information_id"

    def test_header_is_sorted_when_not_given(self) -> None:
        """Without an order, the header is the sorted union of the row keys."""
        document = generate_csv(ROWS).decode(ENCODING)

        assert document.splitlines()[0] == "confidence_score;information_id;provenance;text"

    def test_file_is_utf8_bom_with_crlf(self) -> None:
        """Excel needs the BOM, and RFC 4180 asks for CRLF line endings."""
        content = generate_csv([{"text": "Évry"}])

        assert content.startswith(b"\xef\xbb\xbf")
        assert b"\r\n" in content
        assert "Évry" in content.decode(ENCODING)

    def test_nested_value_keeps_its_information(self) -> None:
        """A nested mapping is written as JSON text, never dropped (§37)."""
        document = generate_csv(ROWS).decode(ENCODING)

        assert '"{\"\"url\"\":\"\"https://exemple.fr\"\"}"' in document

    def test_missing_value_is_an_empty_cell(self) -> None:
        """An absent value stays absent: no placeholder is invented (§24.3)."""
        lines = generate_csv(ROWS).decode(ENCODING).splitlines()
        # Header: confidence_score;information_id;provenance;text
        assert lines[1].split(";")[2] == ""  # row 1 has no provenance
        assert lines[2].split(";")[0] == ""  # row 2 has no confidence

    def test_reproducible(self) -> None:
        """Two exports of the same rows are bit-identical."""
        assert generate_csv(ROWS) == generate_csv(ROWS)

    def test_refuses_a_table_without_columns(self) -> None:
        """A CSV with no column would look like data: it is refused."""
        with pytest.raises(ValidationError, match="at least one column"):
            generate_csv([])
        with pytest.raises(ValidationError, match="at least one column"):
            generate_csv([{}])


class TestJsonExport:
    """§24 delivery — the JSON export is canonical, hence verifiable."""

    def test_bytes_are_exactly_what_the_digest_covers(self) -> None:
        """The client re-computes the ``sha256`` of §24.2 from these bytes."""
        content = generate_json({"b": 1, "a": "é"})

        assert content == b'{"a":"\xc3\xa9","b":1}'

    def test_keys_are_sorted(self) -> None:
        """Deterministic ordering, so two runs of the same payload agree."""
        assert generate_json({"z": 1, "a": 2}) == b'{"a":2,"z":1}'

    def test_refuses_a_null_payload(self) -> None:
        """``null`` is not a delivery."""
        with pytest.raises(ValidationError, match="null payload"):
            generate_json(None)


class TestXmlExport:
    """§24 delivery — the XML export is total and re-parsable."""

    def test_document_is_parsable_and_keeps_the_structure(self) -> None:
        content = generate_xml({"request_id": "REQ_1", "information_units": [{"id": "INF_1"}]})

        root = etree.fromstring(content)
        assert root.tag == "information_package"
        assert root.findtext("request_id") == "REQ_1"
        items = root.find("information_units")
        assert items.get("count") == "1"
        assert items.findtext("item/id") == "INF_1"

    def test_absent_value_is_declared_nil(self) -> None:
        """§37 — absence is stated, never replaced by an empty string."""
        root = etree.fromstring(generate_xml({"summary": None}))

        assert root.find("summary").get("nil") == "true"
        assert root.find("summary").text is None

    def test_keys_are_sorted_and_reproducible(self) -> None:
        first = generate_xml({"b": 1, "a": 2})
        second = generate_xml({"a": 2, "b": 1})

        assert first == second
        assert first.index(b"<a>") < first.index(b"<b>")

    def test_refuses_a_non_mapping_payload(self) -> None:
        with pytest.raises(ValidationError, match="mapping payload"):
            generate_xml(["not", "a", "mapping"])  # type: ignore[arg-type]


class TestXlsxExport:
    """§24.3 — the workbook holds the sheets the delivery can actually fill."""

    def test_sheets_are_the_spec_sheets_in_order(self) -> None:
        workbook = load_workbook(io.BytesIO(generate_xlsx({"data": ROWS})))

        assert [sheet.title for sheet in workbook.worksheets] == [
            SHEET_TITLES[key] for key in SHEET_ORDER
        ]

    def test_data_sheet_holds_the_exported_rows(self) -> None:
        content = generate_xlsx({"data": ROWS}, columns=["information_id", "text"])
        sheet = load_workbook(io.BytesIO(content))[SHEET_TITLES["data"]]

        assert [cell.value for cell in sheet[1]] == ["information_id", "text"]
        assert [cell.value for cell in sheet[2]] == ["INF_1", "Paris"]

    def test_absent_dimension_is_an_empty_sheet_not_an_invented_one(self) -> None:
        """A dimension with no data says so; it does not gain a made-up row."""
        workbook = load_workbook(io.BytesIO(generate_xlsx({"data": ROWS})))
        quality = workbook[SHEET_TITLES["quality"]]

        assert quality.max_row == 1
        assert quality.cell(row=1, column=1).value == "(aucune donnée)"

    def test_content_is_stable_across_generations(self) -> None:
        """Two workbooks carry the same cells (ZIP metadata alone may differ)."""
        first = load_workbook(io.BytesIO(generate_xlsx({"data": ROWS})))
        second = load_workbook(io.BytesIO(generate_xlsx({"data": ROWS})))

        for sheet in first.worksheets:
            rows = [[cell.value for cell in row] for row in sheet.rows]
            other = [[cell.value for cell in row] for row in second[sheet.title].rows]
            assert rows == other


class TestDispatch:
    """A single entry point refuses what it cannot produce (§25.2, §37)."""

    def test_pdf_is_not_available_in_the_v1_stack(self) -> None:
        assert pdf_supported() is False
        assert "pdf" not in AVAILABLE_FORMATS
        with pytest.raises(ValidationError, match="PDF"):
            generate_artifact("pdf", file_stem="REQ_1", payload={})

    def test_evidence_package_is_not_a_file(self) -> None:
        with pytest.raises(ValidationError, match="not a file format"):
            generate_artifact("evidence_package", file_stem="REQ_1", payload={})

    def test_unknown_format_lists_the_known_ones(self) -> None:
        with pytest.raises(ValidationError) as failure:
            generate_artifact("docx", file_stem="REQ_1")

        for known in FILE_FORMATS:
            assert known in str(failure.value)

    def test_file_name_cannot_escape_its_directory(self) -> None:
        generated = generate_artifact("csv", file_stem="../../etc/passwd", rows=[{"a": 1}])

        assert "/" not in generated.file_name
        assert generated.file_name.endswith(".csv")


class TestNothingIsWrittenToDisk:
    """§4 — the generators are pure: bytes in memory, nothing else.

    The delivery path stores the bytes in S3 (§4.3); a generator that wrote a
    temporary file would leave a copy of delivered content on the container's
    filesystem, outside any retention rule. This is what the plan calls the
    "no leak" property — proven here rather than assumed, by generating every
    available format with the working directory pointed at an empty temporary
    folder and checking that it is still empty afterwards.
    """

    @pytest.mark.parametrize("fmt", sorted(AVAILABLE_FORMATS))
    def test_no_file_appears_in_the_working_directory(
        self, fmt: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Generating every available format leaves the directory untouched."""
        monkeypatch.chdir(tmp_path)

        generated = generate_artifact(fmt, file_stem="REQ_LEAK", rows=ROWS)

        assert generated.content, "un générateur rend des octets, jamais un chemin"
        assert list(tmp_path.iterdir()) == []

    def test_content_is_bytes_and_not_a_path(self) -> None:
        """The public API returns the payload itself, not where it was written."""
        generated = generate_artifact("json", file_stem="REQ_LEAK", payload={"a": 1})

        assert isinstance(generated.content, bytes)
        assert generated.file_name == "REQ_LEAK.json"
