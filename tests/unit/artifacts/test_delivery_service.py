"""§24 — delivering the file a request asked for, and stating what could not be.

``deliver_artifacts`` is the seam between the pipeline and the §24.2 records. The
tests below pin its four outcomes: no file (the §7 default), a produced file, an
explicit refusal (``pdf``, unknown format) and an explicit degradation (no object
storage, no database sequence). All of them must be visible in
``limitations`` — a silent empty ``artifacts`` list is exactly the failure mode
P10 of the conformance plan describes.
"""

from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook

from app.api.v1.requests.schemas import RequiredOutput as ApiRequiredOutput
from app.artifacts.delivery.delivery_service import (
    DATA_COLUMNS,
    IN_PROCESS_SEQUENCE_LIMITATION,
    deliver_artifacts,
    output_format_of,
    tabular_projection,
)
from app.artifacts.generators import PDF_UNAVAILABLE_REASON
from app.artifacts.generators.xlsx_generator import SHEET_TITLES
from app.core.hashing import sha256_hex
from app.domain.entities.artifact import Artifact
from app.domain.value_objects.request_constraints import RequiredOutput

UNITS = [
    {
        "information_id": "INF_1",
        "type": "text",
        "source_id": "SRC_1",
        "content": {"text": "Paris est la capitale."},
        "language": "fr",
        "epistemic_status": "factual",
        "data_stage": "raw",
        "confidence": {"score": 0.8, "not_a_probability": True},
        "quality": {},
        "provenance": {"url": "https://fr.wikipedia.org/wiki/Paris", "retrieved_at": "2026-01-01T00:00:00Z"},
    }
]

DELIVERY = {
    "response_id": "RESP_1",
    "request_id": "REQ_1",
    "status": "completed",
    "summary": "Paris est la capitale.",
    "findings": [{"statement": "Paris"}],
    "information_units": UNITS,
    "evidence": [],
    "sources": [{"source_id": "SRC_1", "name": "Wikipédia", "source_type": "web", "trust_level": 7}],
    "artifacts": [],
    "confidence": {"score": 0.8, "not_a_probability": True, "dimensions": {}},
    "provenance": {"pipeline": "PipelineRunner", "request_id": "REQ_1", "plan_id": "PLAN_1", "steps_executed": 2},
    "timestamps": {"started_at": "2026-01-01T00:00:00Z", "completed_at": "2026-01-01T00:00:10Z"},
}


class RecordingStorage:
    """Object-store double that keeps the bytes it received."""

    def __init__(self, *, fail_with: Exception | None = None) -> None:
        self.uploads: list[tuple[str, bytes, str | None]] = []
        self._fail_with = fail_with

    def upload(self, key: str, data: bytes, content_type: str | None = None) -> str:
        """Record the upload, or raise the configured failure."""
        if self._fail_with is not None:
            raise self._fail_with
        self.uploads.append((key, data, content_type))
        return f"s3://inis-artifacts/{key}"


@pytest.fixture(autouse=True)
def _no_database(monkeypatch: pytest.MonkeyPatch) -> None:
    """Run every test without a database unless it sets one itself."""
    monkeypatch.delenv("INIS_DATABASE_URL", raising=False)


class TestOutputFormat:
    """§7 — the format comes from the request, whatever shape it is held in."""

    def test_default_is_the_evidence_package(self) -> None:
        assert output_format_of(None) == "evidence_package"

    def test_mapping_dataclass_and_api_model_agree(self) -> None:
        assert output_format_of({"format": "csv"}) == "csv"
        assert output_format_of(RequiredOutput(format="xlsx")) == "xlsx"
        assert output_format_of(ApiRequiredOutput(format="json")) == "json"

    def test_unknown_format_is_lowercased_not_guessed(self) -> None:
        assert output_format_of({"format": "DOCX"}) == "docx"


class TestTabularProjection:
    """§24.3 « données demandées » — one row per delivered unit."""

    def test_columns_are_the_documented_ones(self) -> None:
        rows = tabular_projection(UNITS)

        assert list(rows[0]) == list(DATA_COLUMNS)

    def test_values_come_from_the_unit(self) -> None:
        row = tabular_projection(UNITS)[0]

        assert row["information_id"] == "INF_1"
        assert row["text"] == "Paris est la capitale."
        assert row["source_url"] == "https://fr.wikipedia.org/wiki/Paris"
        assert row["confidence_score"] == 0.8

    def test_an_absent_dimension_stays_empty(self) -> None:
        """§37 — a unit without quality data exports an empty cell."""
        row = tabular_projection([{"information_id": "INF_2"}])[0]

        assert row["text"] == ""
        assert row["source_url"] is None
        assert row["confidence_score"] is None


class TestDeliverArtifacts:
    """§24 — what the delivery carries, and what it must admit."""

    async def test_evidence_package_attaches_no_file(self) -> None:
        """§24.1 — the package *is* the response; nothing else is claimed."""
        outcome = await deliver_artifacts(request_id="REQ_1", delivery=DELIVERY)

        assert outcome.artifacts == []
        assert outcome.limitations == []

    async def test_json_delivery_is_traceable_and_verifiable(self) -> None:
        storage = RecordingStorage()
        outcome = await deliver_artifacts(
            request_id="REQ_1",
            required_output={"format": "json"},
            delivery=DELIVERY,
            information_units=UNITS,
            storage=storage,
        )

        key, content, content_type = storage.uploads[0]
        record = outcome.artifacts[0]
        assert key.endswith("/REQ_1.json")
        assert content_type == "application/json"
        assert record["artifact_id"].startswith("ART_")
        assert record["request_id"] == "REQ_1"
        assert record["created_at"]
        assert record["file_name"] == "REQ_1.json"
        assert record["sha256"] == sha256_hex(content)
        assert record["size_bytes"] == len(content)
        assert record["source_ids"] == ["SRC_1"]
        assert record["provenance_complete"] is True
        assert record["storage_ref"].startswith("s3://")
        assert outcome.stored is True
        # The only additions to §24.2 are the two fields this endpoint needs.
        assert set(record) - {"request_id", "created_at"} == set(Artifact.model_fields)

    async def test_a_file_does_not_contain_its_own_record(self) -> None:
        """The §24.2 record is published in §24.1, not inside the delivered file."""
        storage = RecordingStorage()
        await deliver_artifacts(
            request_id="REQ_1",
            required_output={"format": "json"},
            delivery=DELIVERY,
            information_units=UNITS,
            storage=storage,
        )

        written = json.loads(storage.uploads[0][1].decode("utf-8"))
        assert "artifacts" not in written
        assert written["request_id"] == "REQ_1"
        assert written["information_units"] == UNITS

    async def test_xlsx_delivery_fills_the_spec_sheets(self) -> None:
        storage = RecordingStorage()
        outcome = await deliver_artifacts(
            request_id="REQ_1",
            required_output={"format": "xlsx"},
            delivery=DELIVERY,
            information_units=UNITS,
            sources=DELIVERY["sources"],
            storage=storage,
        )

        workbook = load_workbook(io.BytesIO(storage.uploads[0][1]))
        assert [sheet.title for sheet in workbook.worksheets] == list(SHEET_TITLES.values())
        assert workbook[SHEET_TITLES["data"]].cell(row=2, column=1).value == "INF_1"
        sources_row = [cell.value for cell in workbook[SHEET_TITLES["sources"]][2]]
        assert "SRC_1" in sources_row and "Wikipédia" in sources_row
        version = workbook[SHEET_TITLES["version"]]
        header = [cell.value for cell in version[1]]
        row = [cell.value for cell in version[2]]
        assert row[header.index("artifact_id")] == outcome.artifacts[0]["artifact_id"]

    async def test_csv_delivery_exports_the_delivered_units(self) -> None:
        storage = RecordingStorage()
        await deliver_artifacts(
            request_id="REQ_1",
            required_output={"format": "csv"},
            delivery=DELIVERY,
            information_units=UNITS,
            storage=storage,
        )

        document = storage.uploads[0][1].decode("utf-8-sig")
        assert document.splitlines()[0].startswith("information_id;")
        assert "Paris est la capitale." in document

    async def test_pdf_is_refused_and_the_refusal_is_stated(self) -> None:
        """§4.1/§37 — a format the stack cannot produce is named, never faked."""
        outcome = await deliver_artifacts(
            request_id="REQ_1",
            required_output={"format": "pdf"},
            delivery=DELIVERY,
            storage=RecordingStorage(),
        )

        assert outcome.artifacts == []
        assert PDF_UNAVAILABLE_REASON in outcome.limitations

    async def test_unknown_format_is_stated_with_the_available_ones(self) -> None:
        outcome = await deliver_artifacts(
            request_id="REQ_1",
            required_output={"format": "docx"},
            delivery=DELIVERY,
            storage=RecordingStorage(),
        )

        assert outcome.artifacts == []
        assert any(
            "docx" in limitation and "csv" in limitation for limitation in outcome.limitations
        )

    async def test_without_object_storage_the_gap_is_explicit(self) -> None:
        outcome = await deliver_artifacts(
            request_id="REQ_1",
            required_output={"format": "csv"},
            delivery=DELIVERY,
            information_units=UNITS,
        )

        assert len(outcome.artifacts) == 1
        assert outcome.stored is False
        assert any("non stocké" in limitation for limitation in outcome.limitations)

    async def test_without_database_the_sequence_is_admitted_as_weak(self) -> None:
        """The §24.2 sequence is only authoritative when it is persisted."""
        outcome = await deliver_artifacts(
            request_id="REQ_1",
            required_output={"format": "csv"},
            delivery=DELIVERY,
            information_units=UNITS,
            storage=RecordingStorage(),
        )

        assert IN_PROCESS_SEQUENCE_LIMITATION in outcome.limitations

    async def test_a_failing_store_does_not_lose_the_record(self) -> None:
        outcome = await deliver_artifacts(
            request_id="REQ_1",
            required_output={"format": "csv"},
            delivery=DELIVERY,
            information_units=UNITS,
            storage=RecordingStorage(fail_with=RuntimeError("S3 unreachable")),
        )

        assert len(outcome.artifacts) == 1
        assert any("S3 unreachable" in limitation for limitation in outcome.limitations)
